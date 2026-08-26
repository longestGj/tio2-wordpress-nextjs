from __future__ import annotations

import argparse
import base64
import os
import shutil
import subprocess
import sys
import tempfile
from copy import deepcopy
from pathlib import Path
from typing import Any

import yaml
from docx import Document
from docx.text.paragraph import Paragraph

from verify_tds import load_manifest, resolve_record_path, sha256, validate_manifest


def _set_paragraph(paragraph: Paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run._element.getparent().remove(run._element)
    else:
        paragraph.add_run(text)


def _set_labeled_paragraph(paragraph: Paragraph, label: str, value: str) -> None:
    if not paragraph.runs:
        paragraph.add_run(label)
        paragraph.add_run(value)
        return
    paragraph.runs[0].text = label
    if len(paragraph.runs) == 1:
        paragraph.add_run(value)
    else:
        paragraph.runs[1].text = value
        for run in paragraph.runs[2:]:
            run._element.getparent().remove(run._element)


def _clone_paragraph(template: Paragraph, parent, text: str) -> Paragraph:
    element = deepcopy(template._p)
    parent._tc.append(element)
    paragraph = Paragraph(element, parent)
    _set_paragraph(paragraph, text)
    return paragraph


def _replace_list_cell(cell, applications: list[str], features: list[str]) -> None:
    paragraphs = list(cell.paragraphs)
    if len(paragraphs) < 9:
        raise ValueError("TEMPLATE_APPLICATION_CELL_INVALID")
    heading_app, item_app, spacer, heading_features, item_feature = (
        paragraphs[0],
        paragraphs[1],
        paragraphs[3],
        paragraphs[4],
        paragraphs[5],
    )
    templates = [deepcopy(paragraph._p) for paragraph in (heading_app, item_app, spacer, heading_features, item_feature)]
    for paragraph in list(cell.paragraphs):
        cell._tc.remove(paragraph._p)

    def append(template_element, text: str) -> None:
        element = deepcopy(template_element)
        cell._tc.append(element)
        _set_paragraph(Paragraph(element, cell), text)

    append(templates[0], "Applications")
    for item in applications:
        append(templates[1], item)
    append(templates[2], "")
    append(templates[3], "Supplier-reported Key Features")
    for item in features:
        append(templates[4], item)


def _technical_table(document: Document):
    for outer in document.tables:
        for row in outer.rows:
            for cell in row.cells:
                for nested in cell.tables:
                    if nested.rows and [c.text.strip() for c in nested.rows[0].cells] == ["Property", "Reported Value"]:
                        return nested
    raise ValueError("TEMPLATE_TECHNICAL_TABLE_MISSING")


def _replace_technical_rows(table, rows: list[dict[str, Any]]) -> None:
    while len(table.rows) - 1 < len(rows):
        table._tbl.append(deepcopy(table.rows[-1]._tr))
    while len(table.rows) - 1 > len(rows):
        table._tbl.remove(table.rows[-1]._tr)
    for index, item in enumerate(rows, start=1):
        _set_paragraph(table.cell(index, 0).paragraphs[0], str(item["property"]))
        _set_paragraph(table.cell(index, 1).paragraphs[0], str(item["value"]))


def build_docx(template: Path, data: dict[str, Any], output: Path) -> None:
    document = Document(template)
    product = data["product"]
    content = data["content"]
    grade = str(product["display_grade"])
    family = str(product["product_family"])
    application_family = str(product["application_family"])

    _set_paragraph(document.paragraphs[0], grade)
    _set_paragraph(document.paragraphs[1], f"{family} | {application_family}".upper())
    _set_paragraph(document.paragraphs[2], str(content["description"]))
    _set_paragraph(document.paragraphs[4], str(content["evidence_note"]))
    status_text = "DRAFT - NOT APPROVED FOR PUBLIC USE" if product["status"] == "draft" else "USER-APPROVED"
    _set_paragraph(document.paragraphs[5], status_text)

    _replace_list_cell(
        document.tables[0].cell(0, 0),
        [str(value) for value in content.get("applications", [])],
        [str(value) for value in content.get("key_features", [])],
    )
    _replace_technical_rows(_technical_table(document), content["technical_data"])

    performance_cell = document.tables[1].cell(0, 0)
    _set_labeled_paragraph(performance_cell.paragraphs[1], "Viscosity  ", str(content["viscosity_note"]))
    _set_labeled_paragraph(performance_cell.paragraphs[2], "Packaging  ", str(content["packaging_note"]))
    safety_cell = document.tables[1].cell(0, 2)
    _set_paragraph(safety_cell.paragraphs[1], str(content["safety_note"]))
    _set_paragraph(document.tables[2].cell(0, 0).paragraphs[1], str(content["release_status"]))

    section = document.sections[0]
    header_cell = section.header.tables[0].cell(0, 1)
    _set_paragraph(header_cell.paragraphs[0], grade)
    _set_paragraph(header_cell.paragraphs[1], family)
    footer_cell = section.footer.tables[0].cell(0, 0)
    _set_paragraph(footer_cell.paragraphs[1], str(content["contact_headline"]))
    _set_paragraph(footer_cell.paragraphs[2], str(content["contact_details"]))
    footer_status = "Draft" if product["status"] == "draft" else "User-approved"
    year = str(product["issue_date"]).split("-", 1)[0]
    _set_paragraph(
        section.footer.paragraphs[1],
        f"TDS | {grade} {family} | {footer_status} | Rev. {product['revision']} | {year}",
    )

    document.core_properties.title = f"{grade} Technical Data Sheet"
    document.core_properties.subject = f"TIOVAR {grade} {product['status']} technical data sheet"
    document.core_properties.comments = "User approval required before public placement" if product["status"] == "draft" else "User-approved"
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)


def _find_renderer() -> Path:
    explicit = os.environ.get("TIOVAR_RENDER_DOCX")
    if explicit:
        path = Path(explicit)
        if path.is_file():
            return path
    base = Path.home() / ".codex" / "plugins" / "cache" / "openai-primary-runtime" / "documents"
    candidates = sorted(base.glob("*/skills/documents/render_docx.py"), reverse=True)
    if not candidates:
        raise FileNotFoundError("DOCUMENT_RENDERER_NOT_FOUND")
    return candidates[0]


def _render_pdf_with_word(docx_path: Path, pdf_path: Path) -> None:
    powershell = shutil.which("powershell.exe") or shutil.which("powershell")
    if not powershell:
        raise FileNotFoundError("POWERSHELL_NOT_FOUND")
    script = r"""
$ErrorActionPreference = 'Stop'
$inputPath = [System.IO.Path]::GetFullPath($env:TIOVAR_INPUT_DOCX)
$outputPath = [System.IO.Path]::GetFullPath($env:TIOVAR_OUTPUT_PDF)
[System.IO.Directory]::CreateDirectory([System.IO.Path]::GetDirectoryName($outputPath)) | Out-Null
$word = $null
$document = $null
try {
    $word = New-Object -ComObject Word.Application
    $word.Visible = $false
    $word.DisplayAlerts = 0
    $document = $word.Documents.Open($inputPath, $false, $true)
    $document.ExportAsFixedFormat($outputPath, 17, $false, 0, 0, 1, 1, 0, $true, $true, 0, $true, $true, $false)
}
finally {
    if ($null -ne $document) { $document.Close($false) }
    if ($null -ne $word) { $word.Quit() }
    if ($null -ne $document) { [void][System.Runtime.InteropServices.Marshal]::ReleaseComObject($document) }
    if ($null -ne $word) { [void][System.Runtime.InteropServices.Marshal]::ReleaseComObject($word) }
    [GC]::Collect()
    [GC]::WaitForPendingFinalizers()
}
"""
    encoded = base64.b64encode(script.encode("utf-16le")).decode("ascii")
    env = os.environ.copy()
    env["TIOVAR_INPUT_DOCX"] = str(docx_path.resolve())
    env["TIOVAR_OUTPUT_PDF"] = str(pdf_path.resolve())
    result = subprocess.run(
        [powershell, "-NoProfile", "-NonInteractive", "-EncodedCommand", encoded],
        env=env,
        capture_output=True,
        text=True,
    )
    if result.returncode != 0 or not pdf_path.is_file():
        raise RuntimeError("WORD_PDF_RENDER_FAILED\n" + result.stdout + result.stderr)


def render_pdf(docx_path: Path, pdf_path: Path) -> None:
    if sys.platform == "win32":
        _render_pdf_with_word(docx_path, pdf_path)
        return
    renderer = _find_renderer()
    with tempfile.TemporaryDirectory(prefix="tiovar-tds-render-") as temp_dir:
        render_dir = Path(temp_dir)
        result = subprocess.run(
            [sys.executable, str(renderer), str(docx_path), "--output_dir", str(render_dir), "--emit_pdf"],
            capture_output=True,
            text=True,
        )
        rendered_pdf = render_dir / f"{docx_path.stem}.pdf"
        if result.returncode != 0 or not rendered_pdf.is_file():
            raise RuntimeError("PDF_RENDER_FAILED\n" + result.stdout + result.stderr)
        pdf_path.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(rendered_pdf, pdf_path)


def _atomic_commit(staged: list[tuple[Path, Path]], staging_dir: Path) -> None:
    backups: dict[Path, Path] = {}
    created_targets: set[Path] = set()
    for _, target in staged:
        if target.is_file():
            backup = staging_dir / f"backup-{len(backups)}-{target.name}"
            shutil.copy2(target, backup)
            backups[target] = backup
        else:
            created_targets.add(target)
    try:
        for source, target in staged:
            os.replace(source, target)
    except Exception:
        for target, backup in backups.items():
            shutil.copy2(backup, target)
        for target in created_targets:
            if target.is_file():
                target.unlink()
        raise


def main() -> int:
    parser = argparse.ArgumentParser(description="Build one TIOVAR English TDS")
    parser.add_argument("--manifest", required=True, type=Path)
    parser.add_argument("--template", required=True, type=Path)
    parser.add_argument("--output-dir", required=True, type=Path)
    parser.add_argument("--docx-only", action="store_true")
    parser.add_argument("--replace-existing-draft", action="store_true")
    args = parser.parse_args()

    try:
        manifest_snapshot = sha256(args.manifest)
        data = load_manifest(args.manifest)
        errors = validate_manifest(data, args.manifest)
        if data.get("product", {}).get("status") != "draft":
            errors.append("BUILD_STATUS_NOT_DRAFT")
        recorded_template = resolve_record_path(
            str(data.get("template", {}).get("path") or ""), args.manifest
        )
        if args.template.resolve() != recorded_template.resolve():
            errors.append("TEMPLATE_ARGUMENT_MISMATCH")
        output_dir = args.output_dir.resolve()
        manifest_dir = args.manifest.resolve().parent
        product_id = str(data.get("product", {}).get("id") or "")
        official_output_dir = (Path(__file__).resolve().parents[3] / "documents" / "tds" / product_id).resolve()
        if output_dir != manifest_dir:
            errors.append("OUTPUT_OUTSIDE_MANIFEST_DIRECTORY")
        if manifest_dir != official_output_dir or output_dir != official_output_dir:
            errors.append("OUTPUT_OUTSIDE_TDS_ROOT")
        if "public" in {part.lower() for part in output_dir.parts}:
            errors.append("OUTPUT_PUBLIC_FORBIDDEN")
        artifacts = data.get("artifacts", {})
        docx_path = output_dir / str(artifacts.get("docx") or "")
        pdf_path = output_dir / str(artifacts.get("pdf") or "")
        target_snapshots = {
            docx_path: sha256(docx_path) if docx_path.is_file() else None,
            pdf_path: sha256(pdf_path) if pdf_path.is_file() else None,
        }
        if not args.replace_existing_draft:
            if docx_path.is_file():
                errors.append("OUTPUT_EXISTS:DOCX")
            if not args.docx_only and pdf_path.is_file():
                errors.append("OUTPUT_EXISTS:PDF")
        else:
            for label, path, hash_field in (
                ("DOCX", docx_path, "docx_sha256"),
                ("PDF", pdf_path, "pdf_sha256"),
            ):
                if label == "PDF" and args.docx_only:
                    continue
                if not path.is_file():
                    continue
                expected_hash = str(artifacts.get(hash_field) or "").lower()
                if len(expected_hash) != 64:
                    errors.append(f"EXISTING_ARTIFACT_HASH_REQUIRED:{label}")
                elif sha256(path) != expected_hash:
                    errors.append(f"EXISTING_ARTIFACT_HASH_MISMATCH:{label}")
        if errors:
            for error in sorted(set(errors)):
                print(error)
            return 1
        if not args.template.is_file():
            print("TEMPLATE_MISSING")
            return 1

        artifacts = data["artifacts"]
        with tempfile.TemporaryDirectory(prefix=".tiovar-tds-stage-", dir=output_dir) as stage_name:
            staging_dir = Path(stage_name)
            staged_docx = staging_dir / docx_path.name
            staged_pdf = staging_dir / pdf_path.name
            staged_manifest = staging_dir / "sources.yaml"
            build_docx(args.template, data, staged_docx)
            artifacts["docx_sha256"] = sha256(staged_docx)
            staged_files = [(staged_docx, docx_path)]
            if not args.docx_only:
                render_pdf(staged_docx, staged_pdf)
                artifacts["pdf_sha256"] = sha256(staged_pdf)
                staged_files.append((staged_pdf, pdf_path))
            data.setdefault("qa", {})["visual_review"] = {
                "status": "pending",
                "page_count": None,
                "docx_sha256": None,
                "pdf_sha256": None,
            }
            staged_manifest.write_text(
                yaml.safe_dump(data, sort_keys=False, allow_unicode=True),
                encoding="utf-8",
            )
            staged_files.append((staged_manifest, args.manifest.resolve()))
            if sha256(args.manifest) != manifest_snapshot:
                raise RuntimeError("CONCURRENT_MODIFICATION:MANIFEST")
            for path, initial_hash in target_snapshots.items():
                current_hash = sha256(path) if path.is_file() else None
                if current_hash != initial_hash:
                    raise RuntimeError(f"CONCURRENT_MODIFICATION:{path.name}")
            _atomic_commit(staged_files, staging_dir)
        print(f"DOCX={docx_path}")
        if not args.docx_only:
            print(f"PDF={pdf_path}")
        return 0
    except Exception as exc:
        print(f"BUILD_FAILED:{type(exc).__name__}:{exc}")
        return 1


if __name__ == "__main__":
    sys.exit(main())
