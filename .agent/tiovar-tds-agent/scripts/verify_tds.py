from __future__ import annotations

import argparse
import hashlib
import re
import sys
import zipfile
from collections import Counter
from pathlib import Path
from typing import Any
from xml.etree import ElementTree

import yaml
from docx import Document
from pypdf import PdfReader


SCHEMA_VERSION = "TIOVAR-TDS-SOURCES-V0.1"
SAFE_ID = re.compile(r"^[a-z0-9]+(?:-[a-z0-9]+)*$")
PROHIBITED_PATTERNS = {
    "equivalent": re.compile(r"\bequivalent\b", re.IGNORECASE),
    "drop-in replacement": re.compile(r"\bdrop[ -]in replacement\b", re.IGNORECASE),
    "Proven Performance": re.compile(r"\bproven performance\b", re.IGNORECASE),
    "registered mark": re.compile(r"®|\bregistered trademark\b", re.IGNORECASE),
    "certification": re.compile(r"\b(?:iso(?:\s*\d+)?\s+certified|certified product|product certification)\b", re.IGNORECASE),
    "regulatory compliance": re.compile(r"\b(?:(?:reach|rohs|fda)\s+(?:compliant|approved|registered)|complies with)\b", re.IGNORECASE),
    "self-manufacturing": re.compile(r"\b(?:(?:we|tiovar)\s+(?:manufacture|manufactures|produce|produces|make|makes)|manufactured by tiovar)\b", re.IGNORECASE),
}
REPO_ROOT = Path(__file__).resolve().parents[3]


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def load_manifest(path: Path) -> dict[str, Any]:
    data = yaml.safe_load(path.read_text(encoding="utf-8"))
    if not isinstance(data, dict):
        raise ValueError("MANIFEST_NOT_MAPPING")
    return data


def resolve_record_path(value: str, manifest_path: Path | None) -> Path:
    path = Path(value)
    if path.is_absolute():
        return path
    if value.replace("\\", "/").startswith(".agent/"):
        return (REPO_ROOT / path).resolve()
    if manifest_path is not None:
        return (manifest_path.resolve().parent / path).resolve()
    return path.resolve()


def _mapping_records(path: Path) -> list[str]:
    if path.suffix.lower() != ".xlsx":
        return [line for line in path.read_text(encoding="utf-8", errors="ignore").splitlines() if line.strip()]
    namespace = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    with zipfile.ZipFile(path) as package:
        shared: list[str] = []
        if "xl/sharedStrings.xml" in package.namelist():
            root = ElementTree.fromstring(package.read("xl/sharedStrings.xml"))
            shared = ["".join(node.text or "" for node in item.findall(".//m:t", namespace)) for item in root.findall("m:si", namespace)]
        records: list[str] = []
        for name in sorted(
            item for item in package.namelist() if item.startswith("xl/worksheets/sheet") and item.endswith(".xml")
        ):
            root = ElementTree.fromstring(package.read(name))
            for row in root.findall(".//m:row", namespace):
                values: list[str] = []
                for cell in row.findall("m:c", namespace):
                    cell_type = cell.get("t")
                    if cell_type == "inlineStr":
                        value = "".join(node.text or "" for node in cell.findall(".//m:t", namespace))
                    else:
                        value_node = cell.find("m:v", namespace)
                        raw = "" if value_node is None else value_node.text or ""
                        value = shared[int(raw)] if cell_type == "s" and raw.isdigit() else raw
                    if value:
                        values.append(value)
                if values:
                    records.append(" | ".join(values))
        return records


def _flatten(value: Any) -> str:
    if isinstance(value, dict):
        return "\n".join(f"{key}: {_flatten(item)}" for key, item in value.items())
    if isinstance(value, list):
        return "\n".join(_flatten(item) for item in value)
    return "" if value is None else str(value)


def _docx_text(path: Path) -> str:
    document = Document(path)
    chunks = [paragraph.text for paragraph in document.paragraphs]

    def visit_table(table) -> None:
        for row in table.rows:
            for cell in row.cells:
                chunks.extend(paragraph.text for paragraph in cell.paragraphs)
                for nested in cell.tables:
                    visit_table(nested)

    for table in document.tables:
        visit_table(table)
    for section in document.sections:
        chunks.extend(paragraph.text for paragraph in section.header.paragraphs)
        chunks.extend(paragraph.text for paragraph in section.footer.paragraphs)
        for table in section.header.tables:
            visit_table(table)
        for table in section.footer.tables:
            visit_table(table)
    return "\n".join(chunks)


def _pdf_text(path: Path) -> tuple[str, int]:
    reader = PdfReader(str(path))
    return "\n".join(page.extract_text() or "" for page in reader.pages), len(reader.pages)


def _normalize_text(value: str) -> str:
    return " ".join(value.replace("\u00a0", " ").split()).casefold()


def _match_tokens(value: str) -> list[str]:
    return re.findall(r"[a-z0-9]+(?:\.[0-9]+)?", value.casefold())


def _contains_sequence(tokens: list[str], sequence: list[str]) -> bool:
    return bool(sequence) and any(tokens[index : index + len(sequence)] == sequence for index in range(len(tokens) - len(sequence) + 1))


def _contains_property_tokens(container_tokens: list[str], property_tokens: list[str]) -> bool:
    if _contains_sequence(container_tokens, property_tokens):
        return True
    return bool(property_tokens) and "".join(property_tokens) in "".join(container_tokens)


def _technical_pair_present(source_text: str, property_name: str, value: str) -> bool:
    source_tokens = _match_tokens(source_text)
    property_tokens = _match_tokens(property_name.split(",", 1)[0])
    value_tokens = _match_tokens(value)
    if not property_tokens or not value_tokens:
        return False
    for index in range(len(source_tokens) - len(property_tokens) + 1):
        if source_tokens[index : index + len(property_tokens)] != property_tokens:
            continue
        window_start = index + len(property_tokens)
        window = source_tokens[window_start : window_start + 7]
        if _contains_sequence(window, value_tokens):
            return True
    source_compact = "".join(source_tokens)
    property_compact = "".join(property_tokens)
    value_compact = "".join(value_tokens)
    property_index = source_compact.find(property_compact)
    if property_index >= 0:
        tail = source_compact[property_index + len(property_compact) : property_index + len(property_compact) + 32]
        if value_compact in tail:
            return True
    return False


def _artifact_tokens(value: str) -> Counter[str]:
    return Counter(_match_tokens(value))


def _expected_artifact_text(data: dict[str, Any]) -> str:
    product = data["product"]
    content = data["content"]
    grade = str(product["display_grade"])
    family = str(product["product_family"])
    application_family = str(product["application_family"])
    status_body = "DRAFT - NOT APPROVED FOR PUBLIC USE" if product["status"] == "draft" else "USER-APPROVED"
    footer_status = "Draft" if product["status"] == "draft" else "User-approved"
    year = str(product["issue_date"]).split("-", 1)[0]
    chunks = [
        grade,
        f"{family} | {application_family}".upper(),
        str(content["description"]),
        str(content["evidence_note"]),
        status_body,
        "Applications",
        *[str(value) for value in content.get("applications", [])],
        "Supplier-reported Key Features",
        *[str(value) for value in content.get("key_features", [])],
        "Reported Technical Data",
        "Property",
        "Reported Value",
    ]
    for row in content.get("technical_data", []):
        chunks.extend([str(row["property"]), str(row["value"])])
    chunks.extend(
        [
            "Performance & Packaging",
            f"Viscosity {content['viscosity_note']}",
            f"Packaging {content['packaging_note']}",
            "Safety / SDS Status",
            str(content["safety_note"]),
            "Release Status",
            str(content["release_status"]),
            f"TDS | {grade} {family} | {footer_status} | Rev. {product['revision']} | {year}",
            grade,
            family,
            "CONTACT",
            str(content["contact_headline"]),
            str(content["contact_details"]),
        ]
    )
    return "\n".join(chunks)


def _required_content_tokens(data: dict[str, Any]) -> list[tuple[str, str]]:
    product = data.get("product", {})
    content = data.get("content", {})
    tokens = [
        ("display_grade", str(product.get("display_grade") or "")),
        ("product_family", str(product.get("product_family") or "")),
        ("application_family", str(product.get("application_family") or "")),
        ("revision", str(product.get("revision") or "")),
        ("description", str(content.get("description") or "")),
        ("viscosity_note", str(content.get("viscosity_note") or "")),
        ("packaging_note", str(content.get("packaging_note") or "")),
        ("safety_note", str(content.get("safety_note") or "")),
        ("evidence_note", str(content.get("evidence_note") or "")),
        ("release_status", str(content.get("release_status") or "")),
        ("contact_headline", str(content.get("contact_headline") or "")),
        ("contact_details", str(content.get("contact_details") or "")),
    ]
    tokens.extend(("application", str(value)) for value in content.get("applications", []))
    tokens.extend(("key_feature", str(value)) for value in content.get("key_features", []))
    for row in content.get("technical_data", []):
        tokens.append(("technical_data", str(row.get("property") or "")))
        tokens.append(("technical_data", str(row.get("value") or "")))
    return [(field, text) for field, text in tokens if text]


def validate_manifest(data: dict[str, Any], manifest_path: Path | None = None) -> list[str]:
    errors: list[str] = []
    if data.get("schema_version") != SCHEMA_VERSION:
        errors.append("SCHEMA_VERSION_INVALID")

    product = data.get("product") if isinstance(data.get("product"), dict) else {}
    if product.get("brand") != "TIOVAR":
        errors.append("WRONG_BRAND")
    if product.get("language") != "en":
        errors.append("WRONG_LANGUAGE")
    product_id = str(product.get("id") or "")
    if not SAFE_ID.fullmatch(product_id):
        errors.append("PRODUCT_ID_INVALID")
    if product.get("status") not in {"draft", "user-approved"}:
        errors.append("STATUS_INVALID")
    for field in ("display_grade", "supplier_grade", "product_family", "application_family", "revision", "issue_date"):
        if not str(product.get(field) or "").strip() or product.get(field) == "REQUIRED":
            errors.append(f"PRODUCT_FIELD_MISSING:{field}")

    sources = data.get("source_files") if isinstance(data.get("source_files"), list) else []
    if not sources:
        errors.append("SOURCE_FILES_MISSING")
    source_ids: set[str] = set()
    source_roles: dict[str, str] = {}
    source_paths: dict[str, Path] = {}
    source_hashes: dict[str, str] = {}
    source_texts: dict[str, str] = {}
    for index, source in enumerate(sources):
        if not isinstance(source, dict):
            errors.append(f"SOURCE_INVALID:{index}")
            continue
        source_id = str(source.get("id") or "")
        if not source_id or source_id in source_ids:
            errors.append(f"SOURCE_ID_INVALID:{index}")
        source_ids.add(source_id)
        role = str(source.get("role") or "")
        if not role:
            errors.append(f"SOURCE_ROLE_MISSING:{source_id}")
        source_roles[source_id] = role
        if not str(source.get("locator") or "").strip():
            errors.append(f"SOURCE_LOCATOR_MISSING:{source_id}")
        source_path = resolve_record_path(str(source.get("path") or ""), manifest_path)
        expected_hash = str(source.get("sha256") or "").lower()
        source_paths[source_id] = source_path
        source_hashes[source_id] = expected_hash
        if not source_path.is_file():
            errors.append(f"SOURCE_MISSING:{source_id}")
        elif not re.fullmatch(r"[0-9a-f]{64}", expected_hash):
            errors.append(f"SOURCE_HASH_INVALID:{source_id}")
        elif sha256(source_path) != expected_hash:
            errors.append(f"SOURCE_HASH_MISMATCH:{source_id}")
        if role == "primary_technical_source" and source_path.is_file():
            try:
                reader = PdfReader(str(source_path))
                if not reader.pages:
                    raise ValueError("empty PDF")
                source_text = "\n".join(page.extract_text() or "" for page in reader.pages)
                source_texts[source_id] = source_text
                if str(product.get("supplier_grade") or "") not in source_text:
                    errors.append(f"PRIMARY_TDS_IDENTITY_MISMATCH:{source_id}")
            except Exception:
                errors.append(f"PRIMARY_TDS_UNREADABLE:{source_id}")
        if role == "internal_grade_mapping" and source_path.is_file():
            try:
                identities = (str(product.get("display_grade") or ""), str(product.get("supplier_grade") or ""))
                normalized_identities = tuple(_normalize_text(identity) for identity in identities)
                records = [_normalize_text(record) for record in _mapping_records(source_path)]
                if not any(all(identity in record for identity in normalized_identities) for record in records):
                    errors.append(f"PRODUCT_MAPPING_PAIR_MISMATCH:{source_id}")
            except Exception:
                errors.append(f"PRODUCT_MAPPING_SOURCE_UNREADABLE:{source_id}")
    if "primary_technical_source" not in source_roles.values():
        errors.append("PRIMARY_TECHNICAL_SOURCE_MISSING")
    if "internal_grade_mapping" not in source_roles.values():
        errors.append("PRODUCT_MAPPING_SOURCE_MISSING")

    template = data.get("template") if isinstance(data.get("template"), dict) else {}
    template_path = resolve_record_path(str(template.get("path") or ""), manifest_path)
    template_hash = str(template.get("sha256") or "").lower()
    if not template_path.is_file():
        errors.append("TEMPLATE_MISSING")
    elif len(template_hash) != 64:
        errors.append("TEMPLATE_HASH_INVALID")
    elif sha256(template_path) != template_hash:
        errors.append("TEMPLATE_HASH_MISMATCH")
    logo_source_id = str(template.get("embedded_logo_source_id") or "")
    embedded_logo_path = str(template.get("embedded_logo_path") or "")
    if source_roles.get(logo_source_id) != "primary_logo":
        errors.append("TEMPLATE_LOGO_SOURCE_INVALID")
    elif template_path.is_file() and embedded_logo_path:
        try:
            with zipfile.ZipFile(template_path) as package:
                embedded_logo_hash = hashlib.sha256(package.read(embedded_logo_path)).hexdigest()
            if embedded_logo_hash != source_hashes.get(logo_source_id):
                errors.append("TEMPLATE_LOGO_HASH_MISMATCH")
        except Exception:
            errors.append("TEMPLATE_LOGO_UNREADABLE")
    else:
        errors.append("TEMPLATE_LOGO_PATH_MISSING")

    content = data.get("content") if isinstance(data.get("content"), dict) else {}
    if not str(content.get("description") or "").strip() or content.get("description") == "REQUIRED":
        errors.append("DESCRIPTION_MISSING")
    technical_data = content.get("technical_data") if isinstance(content.get("technical_data"), list) else []
    if not technical_data:
        errors.append("TECHNICAL_DATA_MISSING")
    for index, row in enumerate(technical_data):
        if not isinstance(row, dict):
            errors.append(f"TECHNICAL_ROW_INVALID:{index}")
            continue
        for field in ("property", "value", "source_id", "source_locator"):
            if not str(row.get(field) or "").strip() or row.get(field) == "REQUIRED":
                errors.append(f"TECHNICAL_FIELD_MISSING:{index}:{field}")
        if row.get("source_id") not in source_ids:
            errors.append(f"TECHNICAL_SOURCE_UNKNOWN:{index}")
        elif source_roles.get(str(row.get("source_id"))) != "primary_technical_source":
            errors.append(f"TECHNICAL_SOURCE_NOT_PRIMARY:{index}")
        else:
            property_name = str(row.get("property") or "")
            value = str(row.get("value") or "")
            locator = str(row.get("source_locator") or "")
            property_tokens = _match_tokens(property_name.split(",", 1)[0])
            if not _contains_property_tokens(_match_tokens(locator), property_tokens):
                errors.append(f"TECHNICAL_LOCATOR_PROPERTY_MISMATCH:{index}")
            if not _technical_pair_present(source_texts.get(str(row.get("source_id")), ""), property_name, value):
                errors.append(f"TECHNICAL_PROPERTY_VALUE_PAIR_NOT_FOUND:{index}")

    claims = data.get("claims") if isinstance(data.get("claims"), list) else []
    claim_keys: set[tuple[str, str]] = set()
    for index, claim in enumerate(claims):
        if not isinstance(claim, dict):
            errors.append(f"CLAIM_INVALID:{index}")
            continue
        field = str(claim.get("field") or "")
        text = str(claim.get("text") or "")
        support_type = str(claim.get("support_type") or "")
        source_id = str(claim.get("source_id") or "")
        locator = str(claim.get("source_locator") or "")
        source_excerpt = str(claim.get("source_excerpt") or "")
        if not field or not text or support_type not in {"source_supported", "administrative_hold"}:
            errors.append(f"CLAIM_FIELD_MISSING:{index}")
            continue
        claim_keys.add((field, text))
        if support_type == "source_supported":
            if not source_id or not locator or not source_excerpt:
                errors.append(f"CLAIM_SOURCE_FIELDS_MISSING:{index}")
            elif source_roles.get(source_id) != "primary_technical_source":
                errors.append(f"CLAIM_SOURCE_NOT_PRIMARY:{index}")
            elif _normalize_text(source_excerpt) not in _normalize_text(source_texts.get(source_id, "")):
                errors.append(f"CLAIM_SOURCE_EXCERPT_NOT_FOUND:{index}")
        else:
            if field not in {"packaging_note", "safety_note", "release_status", "contact_headline", "contact_details"}:
                errors.append(f"ADMINISTRATIVE_HOLD_FIELD_INVALID:{index}")
            if source_id or locator or source_excerpt:
                errors.append(f"ADMINISTRATIVE_HOLD_SOURCE_FORBIDDEN:{index}")
            hold_language = re.search(
                r"\b(?:pending|unconfirmed|withheld|required|concept|not\s+(?:provided|stated|confirmed|approved))\b",
                text,
                re.IGNORECASE,
            )
            assertive_language = re.search(
                r"\b(?:improves?|increases?|reduces?|ensures?|provides?|contains?|meets?|passes?|complies?)\b|"
                r"\b\d+(?:\.\d+)?\s*(?:kg|g|lb|oz|l|ml|%|mpa|cps|pa\.s)\b",
                text,
                re.IGNORECASE,
            )
            if not hold_language or assertive_language:
                errors.append(f"ADMINISTRATIVE_HOLD_TEXT_UNSAFE:{index}")
    expected_claims = [("description", str(content.get("description") or ""))]
    expected_claims.extend(("application", str(value)) for value in content.get("applications", []))
    expected_claims.extend(("key_feature", str(value)) for value in content.get("key_features", []))
    expected_claims.extend(
        (field, str(content.get(field) or ""))
        for field in (
            "viscosity_note",
            "packaging_note",
            "safety_note",
            "evidence_note",
            "release_status",
            "contact_headline",
            "contact_details",
        )
    )
    for field, text in expected_claims:
        if (field, text) not in claim_keys:
            errors.append(f"CLAIM_SOURCE_MISSING:{field}:{text}")

    artifacts = data.get("artifacts") if isinstance(data.get("artifacts"), dict) else {}
    expected_docx = f"tiovar-{product_id}-en.docx"
    expected_pdf = f"tiovar-{product_id}-en.pdf"
    if artifacts.get("docx") != expected_docx:
        errors.append("ARTIFACT_NAME_INVALID:DOCX")
    if artifacts.get("pdf") != expected_pdf:
        errors.append("ARTIFACT_NAME_INVALID:PDF")

    conflicts = data.get("conflicts")
    if not isinstance(conflicts, list):
        errors.append("CONFLICT_REGISTER_MISSING")
    else:
        for index, conflict in enumerate(conflicts):
            if not isinstance(conflict, dict):
                errors.append(f"SOURCE_CONFLICT:INDEX-{index}")
                continue
            if conflict.get("status") != "resolved":
                errors.append(f"SOURCE_CONFLICT:{conflict.get('id') or f'INDEX-{index}'}")

    manifest_text = _flatten(data)
    for label, pattern in PROHIBITED_PATTERNS.items():
        if pattern.search(manifest_text):
            errors.append(f"PROHIBITED_TERM:{label}")
    return errors


def verify_artifacts(
    data: dict[str, Any], manifest_path: Path, docx_path: Path, pdf_path: Path
) -> tuple[list[str], int]:
    errors = validate_manifest(data, manifest_path)
    page_count = 0
    product = data.get("product", {})
    grade = str(product.get("display_grade") or "")
    artifacts = data.get("artifacts") if isinstance(data.get("artifacts"), dict) else {}
    expected_docx_path = manifest_path.resolve().parent / str(artifacts.get("docx") or "")
    expected_pdf_path = manifest_path.resolve().parent / str(artifacts.get("pdf") or "")
    if docx_path.resolve() != expected_docx_path.resolve():
        errors.append("ARTIFACT_PATH_MISMATCH:DOCX")
    if pdf_path.resolve() != expected_pdf_path.resolve():
        errors.append("ARTIFACT_PATH_MISMATCH:PDF")

    for label, path in (("DOCX", docx_path), ("PDF", pdf_path)):
        if not path.is_file():
            errors.append(f"{label}_MISSING")
        if product.get("status") == "draft" and "public" in {part.lower() for part in path.parts}:
            errors.append(f"DRAFT_IN_PUBLIC:{label}")

    if docx_path.is_file():
        text = _docx_text(docx_path)
        normalized_docx = _normalize_text(text)
        if grade not in text:
            errors.append("DOCX_GRADE_MISSING")
        for label, pattern in PROHIBITED_PATTERNS.items():
            if pattern.search(text):
                errors.append(f"DOCX_PROHIBITED_TERM:{label}")
        expected = artifacts.get("docx_sha256")
        if not expected:
            errors.append("ARTIFACT_HASH_REQUIRED:DOCX")
        elif sha256(docx_path) != str(expected).lower():
            errors.append("DOCX_HASH_MISMATCH")
        for field, token in _required_content_tokens(data):
            if _normalize_text(token) not in normalized_docx:
                errors.append(f"DOCX_CONTENT_MISSING:{field}:{token}")
        if _artifact_tokens(text) != _artifact_tokens(_expected_artifact_text(data)):
            errors.append("DOCX_UNEXPECTED_CONTENT")

    if pdf_path.is_file():
        try:
            text, page_count = _pdf_text(pdf_path)
            normalized_pdf = _normalize_text(text)
            if page_count < 1:
                errors.append("PDF_EMPTY")
            if grade not in text:
                errors.append("PDF_GRADE_MISSING")
            for label, pattern in PROHIBITED_PATTERNS.items():
                if pattern.search(text):
                    errors.append(f"PDF_PROHIBITED_TERM:{label}")
            for field, token in _required_content_tokens(data):
                if _normalize_text(token) not in normalized_pdf:
                    errors.append(f"PDF_CONTENT_MISSING:{field}:{token}")
            if _artifact_tokens(text) != _artifact_tokens(_expected_artifact_text(data)):
                errors.append("PDF_UNEXPECTED_CONTENT")
        except Exception as exc:  # pragma: no cover - exact parser errors vary
            errors.append(f"PDF_INVALID:{type(exc).__name__}")
        expected = artifacts.get("pdf_sha256")
        if not expected:
            errors.append("ARTIFACT_HASH_REQUIRED:PDF")
        elif sha256(pdf_path) != str(expected).lower():
            errors.append("PDF_HASH_MISMATCH")

    visual_review = data.get("qa", {}).get("visual_review", {})
    if visual_review.get("status") != "passed":
        errors.append("VISUAL_REVIEW_NOT_PASSED")
    else:
        if visual_review.get("docx_sha256") != artifacts.get("docx_sha256"):
            errors.append("VISUAL_REVIEW_HASH_MISMATCH:DOCX")
        if visual_review.get("pdf_sha256") != artifacts.get("pdf_sha256"):
            errors.append("VISUAL_REVIEW_HASH_MISMATCH:PDF")
        if visual_review.get("page_count") != page_count:
            errors.append("VISUAL_REVIEW_PAGE_COUNT_MISMATCH")

    return errors, page_count


def main() -> int:
    parser = argparse.ArgumentParser(description="Verify a TIOVAR English TDS package")
    parser.add_argument("--manifest", required=True, type=Path)
    parser.add_argument("--docx", type=Path)
    parser.add_argument("--pdf", type=Path)
    parser.add_argument("--manifest-only", action="store_true")
    args = parser.parse_args()

    try:
        data = load_manifest(args.manifest)
    except Exception as exc:
        print(f"MANIFEST_READ_FAILED:{type(exc).__name__}")
        return 1

    if args.manifest_only:
        errors = validate_manifest(data, args.manifest)
        page_count = 0
    else:
        if not args.docx or not args.pdf:
            print("DOCX_AND_PDF_REQUIRED")
            return 1
        errors, page_count = verify_artifacts(data, args.manifest, args.docx, args.pdf)

    if errors:
        for error in sorted(set(errors)):
            print(error)
        return 1

    technical_rows = len(data.get("content", {}).get("technical_data", []))
    print("TDS_VERIFIED")
    print(f"TECHNICAL_ROWS={technical_rows}")
    if not args.manifest_only:
        print(f"PDF_PAGES={page_count}")
    print("PROHIBITED_TERMS=0")
    print(f"SOURCE_HASHES_MATCH={len(data.get('source_files', []))}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
