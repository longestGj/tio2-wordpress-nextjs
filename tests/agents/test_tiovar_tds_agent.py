from __future__ import annotations

import hashlib
import os
import shutil
import subprocess
import sys
import tempfile
import unittest
import zipfile
from contextlib import contextmanager
from pathlib import Path

import yaml
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject


REPO_ROOT = Path(__file__).resolve().parents[2]
AGENT_ROOT = REPO_ROOT / ".agent" / "tiovar-tds-agent"
PYTHON = Path(sys.executable)


def _all_text(document: Document) -> str:
    chunks = [paragraph.text for paragraph in document.paragraphs]

    def visit_table(table) -> None:
        for row in table.rows:
            for cell in row.cells:
                chunks.extend(paragraph.text for paragraph in cell.paragraphs)
                for nested in cell.tables:
                    visit_table(nested)

    for table in document.tables:
        visit_table(table)
    for section in document.sections:
        chunks.extend(paragraph.text for paragraph in section.header.paragraphs)
        chunks.extend(paragraph.text for paragraph in section.footer.paragraphs)
        for table in section.header.tables:
            visit_table(table)
        for table in section.footer.tables:
            visit_table(table)
    return "\n".join(chunks)


def _write_supplier_pdf(path: Path) -> None:
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})}
    )
    stream = DecodedStreamObject()
    stream.set_data(
        b"BT /F1 12 Tf 72 720 Td "
        b"(SUP-100 Technical Data Sheet) Tj 0 -20 Td "
        b"(TiO2 content, %: 95) Tj 0 -20 Td "
        b"(High hiding power) Tj 0 -20 Td "
        b"(Formulation evaluation in water-based coating) Tj ET"
    )
    page[NameObject("/Contents")] = writer._add_object(stream)
    with path.open("wb") as output:
        writer.write(output)


@contextmanager
def _official_tds_directory():
    root = REPO_ROOT / "documents" / "tds" / "tp-x100"
    if root.exists():
        raise RuntimeError(f"test output already exists: {root}")
    root.mkdir(parents=True)
    try:
        yield str(root)
    finally:
        shutil.rmtree(root)


def _manifest(source_path: Path, *, brand: str = "TIOVAR", language: str = "en") -> dict:
    source_hash = hashlib.sha256(source_path.read_bytes()).hexdigest()
    mapping_path = source_path.with_name("product-mapping.txt")
    mapping_path.write_text("TP-X100 -> SUP-100", encoding="utf-8")
    mapping_hash = hashlib.sha256(mapping_path.read_bytes()).hexdigest()
    template_path = AGENT_ROOT / "templates" / "tiovar-tds-template.docx"
    template_hash = hashlib.sha256(template_path.read_bytes()).hexdigest() if template_path.is_file() else "0" * 64
    logo_path = source_path.with_name("tiovar-logo.png")
    if template_path.is_file():
        with zipfile.ZipFile(template_path) as package:
            logo_path.write_bytes(package.read("word/media/image1.png"))
    else:
        logo_path.write_bytes(b"logo")
    logo_hash = hashlib.sha256(logo_path.read_bytes()).hexdigest()
    return {
        "schema_version": "TIOVAR-TDS-SOURCES-V0.1",
        "product": {
            "id": "tp-x100",
            "brand": brand,
            "display_grade": "TP-X100",
            "supplier_grade": "SUP-100",
            "product_family": "Rutile Titanium Dioxide",
            "application_family": "Water-Based Emulsion Paint",
            "language": language,
            "revision": "0.1",
            "issue_date": "2026-08-26",
            "status": "draft",
        },
        "source_files": [
            {
                "id": "S000",
                "role": "internal_grade_mapping",
                "path": mapping_path.as_posix(),
                "sha256": mapping_hash,
                "locator": "line 1",
            },
            {
                "id": "S001",
                "role": "primary_technical_source",
                "path": source_path.as_posix(),
                "sha256": source_hash,
                "locator": "page 1",
            },
            {
                "id": "S002",
                "role": "primary_logo",
                "path": logo_path.as_posix(),
                "sha256": logo_hash,
                "locator": "full image",
            },
        ],
        "template": {
            "path": ".agent/tiovar-tds-agent/templates/tiovar-tds-template.docx",
            "sha256": template_hash,
            "embedded_logo_source_id": "S002",
            "embedded_logo_path": "word/media/image1.png",
        },
        "content": {
            "description": "TIOVAR TP-X100 is a source-supported rutile titanium dioxide pigment for formulation evaluation.",
            "applications": ["Formulation evaluation in water-based coating"],
            "key_features": ["High hiding power"],
            "technical_data": [
                {
                    "property": "TiO2 content, %",
                    "value": "95",
                    "source_id": "S001",
                    "source_locator": "page 1 / specifications / TiO2 content",
                }
            ],
            "viscosity_note": "No source-supported viscosity result is available.",
            "packaging_note": "Packaging is not confirmed.",
            "safety_note": "An aligned SDS was not provided for this draft.",
            "evidence_note": "Supplier-reported values reproduce the primary supplier TDS.",
            "release_status": "Pending user confirmation before public use.",
            "contact_headline": "Company and contact details pending user approval",
            "contact_details": "Legal name / address / email / website: pending user approval",
        },
        "claims": [
            {
                "id": "C001",
                "field": "description",
                "text": "TIOVAR TP-X100 is a source-supported rutile titanium dioxide pigment for formulation evaluation.",
                "support_type": "source_supported",
                "source_id": "S001",
                "source_locator": "page 1",
                "source_excerpt": "SUP-100 Technical Data Sheet",
            },
            {
                "id": "C002",
                "field": "application",
                "text": "Formulation evaluation in water-based coating",
                "support_type": "source_supported",
                "source_id": "S001",
                "source_locator": "page 1",
                "source_excerpt": "Formulation evaluation in water-based coating",
            },
            {
                "id": "C003",
                "field": "key_feature",
                "text": "High hiding power",
                "support_type": "source_supported",
                "source_id": "S001",
                "source_locator": "page 1",
                "source_excerpt": "High hiding power",
            },
            {
                "id": "C004",
                "field": "viscosity_note",
                "text": "No source-supported viscosity result is available.",
                "support_type": "source_supported",
                "source_id": "S001",
                "source_locator": "page 1",
                "source_excerpt": "SUP-100 Technical Data Sheet",
            },
            {
                "id": "C005",
                "field": "packaging_note",
                "text": "Packaging is not confirmed.",
                "support_type": "administrative_hold",
                "source_id": None,
                "source_locator": None,
                "source_excerpt": None,
            },
            {
                "id": "C006",
                "field": "safety_note",
                "text": "An aligned SDS was not provided for this draft.",
                "support_type": "administrative_hold",
                "source_id": None,
                "source_locator": None,
                "source_excerpt": None,
            },
            {
                "id": "C007",
                "field": "evidence_note",
                "text": "Supplier-reported values reproduce the primary supplier TDS.",
                "support_type": "source_supported",
                "source_id": "S001",
                "source_locator": "page 1",
                "source_excerpt": "Technical Data Sheet",
            },
            {
                "id": "C008",
                "field": "release_status",
                "text": "Pending user confirmation before public use.",
                "support_type": "administrative_hold",
                "source_id": None,
                "source_locator": None,
                "source_excerpt": None,
            },
            {
                "id": "C009",
                "field": "contact_headline",
                "text": "Company and contact details pending user approval",
                "support_type": "administrative_hold",
                "source_id": None,
                "source_locator": None,
                "source_excerpt": None,
            },
            {
                "id": "C010",
                "field": "contact_details",
                "text": "Legal name / address / email / website: pending user approval",
                "support_type": "administrative_hold",
                "source_id": None,
                "source_locator": None,
                "source_excerpt": None,
            },
        ],
        "conflicts": [],
        "unresolved_fields": [],
        "artifacts": {
            "docx": "tiovar-tp-x100-en.docx",
            "pdf": "tiovar-tp-x100-en.pdf",
            "docx_sha256": None,
            "pdf_sha256": None,
        },
        "qa": {
            "visual_review": {
                "status": "pending",
                "page_count": None,
                "docx_sha256": None,
                "pdf_sha256": None,
            }
        },
    }


class TiovarTdsAgentTests(unittest.TestCase):
    def test_build_cli_creates_product_docx_without_leaking_benchmark_identity(self) -> None:
        with _official_tds_directory() as temp_dir:
            root = Path(temp_dir)
            source = root / "supplier.pdf"
            _write_supplier_pdf(source)
            manifest = root / "sources.yaml"
            manifest.write_text(
                yaml.safe_dump(_manifest(source), sort_keys=False, allow_unicode=True),
                encoding="utf-8",
            )
            output_dir = root

            result = subprocess.run(
                [
                    str(PYTHON),
                    str(AGENT_ROOT / "scripts" / "build_tds.py"),
                    "--manifest",
                    str(manifest),
                    "--template",
                    str(AGENT_ROOT / "templates" / "tiovar-tds-template.docx"),
                    "--output-dir",
                    str(output_dir),
                    "--docx-only",
                ],
                cwd=REPO_ROOT,
                capture_output=True,
                text=True,
            )

            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            output = output_dir / "tiovar-tp-x100-en.docx"
            self.assertTrue(output.is_file())
            text = _all_text(Document(output))
            self.assertIn("TP-X100", text)
            self.assertIn("TiO2 content, %", text)
            self.assertNotIn("TP-C120", text)

    def test_build_invalidates_any_previous_visual_review(self) -> None:
        with _official_tds_directory() as temp_dir:
            root = Path(temp_dir)
            source = root / "supplier.pdf"
            _write_supplier_pdf(source)
            data = _manifest(source)
            data["qa"]["visual_review"] = {
                "status": "passed",
                "page_count": 1,
                "docx_sha256": "a" * 64,
                "pdf_sha256": "b" * 64,
            }
            manifest = root / "sources.yaml"
            manifest.write_text(yaml.safe_dump(data, sort_keys=False), encoding="utf-8")

            result = subprocess.run(
                [
                    str(PYTHON),
                    str(AGENT_ROOT / "scripts" / "build_tds.py"),
                    "--manifest",
                    str(manifest),
                    "--template",
                    str(AGENT_ROOT / "templates" / "tiovar-tds-template.docx"),
                    "--output-dir",
                    str(root),
                    "--docx-only",
                ],
                cwd=REPO_ROOT,
                capture_output=True,
                text=True,
            )

            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            updated = yaml.safe_load(manifest.read_text(encoding="utf-8"))
            review = updated["qa"]["visual_review"]
            self.assertEqual(review["status"], "pending")
            self.assertIsNone(review["page_count"])
            self.assertIsNone(review["docx_sha256"])
            self.assertIsNone(review["pdf_sha256"])

    @unittest.skipUnless(sys.platform == "win32", "Word PDF integration is Windows-specific")
    def test_build_cli_creates_a_readable_pdf_with_the_local_renderer(self) -> None:
        with _official_tds_directory() as temp_dir:
            root = Path(temp_dir)
            source = root / "supplier.pdf"
            _write_supplier_pdf(source)
            manifest = root / "sources.yaml"
            manifest.write_text(
                yaml.safe_dump(_manifest(source), sort_keys=False, allow_unicode=True),
                encoding="utf-8",
            )
            output_dir = root

            result = subprocess.run(
                [
                    str(PYTHON),
                    str(AGENT_ROOT / "scripts" / "build_tds.py"),
                    "--manifest",
                    str(manifest),
                    "--template",
                    str(AGENT_ROOT / "templates" / "tiovar-tds-template.docx"),
                    "--output-dir",
                    str(output_dir),
                ],
                cwd=REPO_ROOT,
                capture_output=True,
                text=True,
            )

            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            output = output_dir / "tiovar-tp-x100-en.pdf"
            self.assertTrue(output.is_file())
            from pypdf import PdfReader

            reader = PdfReader(str(output))
            self.assertEqual(len(reader.pages), 1)
            self.assertIn("TP-X100", reader.pages[0].extract_text() or "")

    def test_verify_cli_rejects_prohibited_claim(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            source = root / "supplier.pdf"
            _write_supplier_pdf(source)
            data = _manifest(source)
            data["content"]["description"] += " It is a drop-in replacement."
            manifest = root / "sources.yaml"
            manifest.write_text(yaml.safe_dump(data, sort_keys=False), encoding="utf-8")

            result = subprocess.run(
                [
                    str(PYTHON),
                    str(AGENT_ROOT / "scripts" / "verify_tds.py"),
                    "--manifest",
                    str(manifest),
                    "--manifest-only",
                ],
                cwd=REPO_ROOT,
                capture_output=True,
                text=True,
            )

            self.assertNotEqual(result.returncode, 0)
            self.assertIn("PROHIBITED_TERM", result.stdout + result.stderr)

    def test_verify_cli_blocks_unsupported_certification_regulatory_and_manufacturing_claims(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            source = root / "supplier.pdf"
            _write_supplier_pdf(source)
            data = _manifest(source)
            data["content"]["description"] = "ISO certified, REACH compliant, and TIOVAR produces this pigment."
            data["claims"][0]["text"] = data["content"]["description"]
            manifest = root / "sources.yaml"
            manifest.write_text(yaml.safe_dump(data, sort_keys=False), encoding="utf-8")

            result = subprocess.run(
                [
                    str(PYTHON),
                    str(AGENT_ROOT / "scripts" / "verify_tds.py"),
                    "--manifest",
                    str(manifest),
                    "--manifest-only",
                ],
                cwd=REPO_ROOT,
                capture_output=True,
                text=True,
            )

            combined = result.stdout + result.stderr
            self.assertNotEqual(result.returncode, 0)
            self.assertIn("PROHIBITED_TERM:certification", combined)
            self.assertIn("PROHIBITED_TERM:regulatory compliance", combined)
            self.assertIn("PROHIBITED_TERM:self-manufacturing", combined)

    def test_verify_cli_rejects_wrong_brand_and_language(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            source = root / "supplier.pdf"
            _write_supplier_pdf(source)
            manifest = root / "sources.yaml"
            manifest.write_text(
                yaml.safe_dump(_manifest(source, brand="OTHER", language="zh"), sort_keys=False),
                encoding="utf-8",
            )

            result = subprocess.run(
                [
                    str(PYTHON),
                    str(AGENT_ROOT / "scripts" / "verify_tds.py"),
                    "--manifest",
                    str(manifest),
                    "--manifest-only",
                ],
                cwd=REPO_ROOT,
                capture_output=True,
                text=True,
            )

            self.assertNotEqual(result.returncode, 0)
            combined = result.stdout + result.stderr
            self.assertIn("WRONG_BRAND", combined)
            self.assertIn("WRONG_LANGUAGE", combined)

    def test_verify_cli_rejects_source_hash_mismatch(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            source = root / "supplier.pdf"
            _write_supplier_pdf(source)
            data = _manifest(source)
            data["source_files"][0]["sha256"] = "0" * 64
            manifest = root / "sources.yaml"
            manifest.write_text(yaml.safe_dump(data, sort_keys=False), encoding="utf-8")

            result = subprocess.run(
                [
                    str(PYTHON),
                    str(AGENT_ROOT / "scripts" / "verify_tds.py"),
                    "--manifest",
                    str(manifest),
                    "--manifest-only",
                ],
                cwd=REPO_ROOT,
                capture_output=True,
                text=True,
            )

            self.assertNotEqual(result.returncode, 0)
            self.assertIn("SOURCE_HASH_MISMATCH", result.stdout + result.stderr)

    def test_verify_cli_requires_primary_technical_and_mapping_sources(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            source = root / "supplier.pdf"
            _write_supplier_pdf(source)
            data = _manifest(source)
            data["source_files"] = [item for item in data["source_files"] if item["role"] == "primary_logo"]
            manifest = root / "sources.yaml"
            manifest.write_text(yaml.safe_dump(data, sort_keys=False), encoding="utf-8")

            result = subprocess.run(
                [
                    str(PYTHON),
                    str(AGENT_ROOT / "scripts" / "verify_tds.py"),
                    "--manifest",
                    str(manifest),
                    "--manifest-only",
                ],
                cwd=REPO_ROOT,
                capture_output=True,
                text=True,
            )

            combined = result.stdout + result.stderr
            self.assertNotEqual(result.returncode, 0)
            self.assertIn("PRIMARY_TECHNICAL_SOURCE_MISSING", combined)
            self.assertIn("PRODUCT_MAPPING_SOURCE_MISSING", combined)

    def test_verify_cli_rejects_technical_data_citing_a_visual_source(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            source = root / "supplier.pdf"
            _write_supplier_pdf(source)
            data = _manifest(source)
            data["content"]["technical_data"][0]["source_id"] = "S002"
            manifest = root / "sources.yaml"
            manifest.write_text(yaml.safe_dump(data, sort_keys=False), encoding="utf-8")

            result = subprocess.run(
                [
                    str(PYTHON),
                    str(AGENT_ROOT / "scripts" / "verify_tds.py"),
                    "--manifest",
                    str(manifest),
                    "--manifest-only",
                ],
                cwd=REPO_ROOT,
                capture_output=True,
                text=True,
            )

            self.assertNotEqual(result.returncode, 0)
            self.assertIn("TECHNICAL_SOURCE_NOT_PRIMARY:0", result.stdout + result.stderr)

    def test_verify_cli_rejects_a_technical_value_absent_from_the_primary_tds(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            source = root / "supplier.pdf"
            _write_supplier_pdf(source)
            data = _manifest(source)
            data["content"]["technical_data"][0]["value"] = "999"
            manifest = root / "sources.yaml"
            manifest.write_text(yaml.safe_dump(data, sort_keys=False), encoding="utf-8")

            result = subprocess.run(
                [
                    str(PYTHON),
                    str(AGENT_ROOT / "scripts" / "verify_tds.py"),
                    "--manifest",
                    str(manifest),
                    "--manifest-only",
                ],
                cwd=REPO_ROOT,
                capture_output=True,
                text=True,
            )

            self.assertNotEqual(result.returncode, 0)
            self.assertIn("TECHNICAL_PROPERTY_VALUE_PAIR_NOT_FOUND:0", result.stdout + result.stderr)

    def test_verify_cli_holds_an_unresolved_source_conflict(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            source = root / "supplier.pdf"
            _write_supplier_pdf(source)
            data = _manifest(source)
            data["conflicts"] = [
                {
                    "id": "X001",
                    "field": "TiO2 content",
                    "status": "unresolved",
                    "detail": "Two sources disagree.",
                }
            ]
            manifest = root / "sources.yaml"
            manifest.write_text(yaml.safe_dump(data, sort_keys=False), encoding="utf-8")

            result = subprocess.run(
                [
                    str(PYTHON),
                    str(AGENT_ROOT / "scripts" / "verify_tds.py"),
                    "--manifest",
                    str(manifest),
                    "--manifest-only",
                ],
                cwd=REPO_ROOT,
                capture_output=True,
                text=True,
            )

            self.assertNotEqual(result.returncode, 0)
            self.assertIn("SOURCE_CONFLICT:X001", result.stdout + result.stderr)

    def test_verify_cli_requires_source_mapping_for_editorial_claims(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            source = root / "supplier.pdf"
            _write_supplier_pdf(source)
            data = _manifest(source)
            data["claims"] = [claim for claim in data["claims"] if claim["field"] != "key_feature"]
            manifest = root / "sources.yaml"
            manifest.write_text(yaml.safe_dump(data, sort_keys=False), encoding="utf-8")

            result = subprocess.run(
                [
                    str(PYTHON),
                    str(AGENT_ROOT / "scripts" / "verify_tds.py"),
                    "--manifest",
                    str(manifest),
                    "--manifest-only",
                ],
                cwd=REPO_ROOT,
                capture_output=True,
                text=True,
            )

            self.assertNotEqual(result.returncode, 0)
            self.assertIn("CLAIM_SOURCE_MISSING:key_feature:High hiding power", result.stdout + result.stderr)

    def test_verify_cli_requires_a_claim_excerpt_present_in_the_primary_tds(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            source = root / "supplier.pdf"
            _write_supplier_pdf(source)
            data = _manifest(source)
            data["claims"][0]["source_excerpt"] = "This text is absent from the supplier TDS"
            manifest = root / "sources.yaml"
            manifest.write_text(yaml.safe_dump(data, sort_keys=False), encoding="utf-8")

            result = subprocess.run(
                [
                    str(PYTHON),
                    str(AGENT_ROOT / "scripts" / "verify_tds.py"),
                    "--manifest",
                    str(manifest),
                    "--manifest-only",
                ],
                cwd=REPO_ROOT,
                capture_output=True,
                text=True,
            )

            self.assertNotEqual(result.returncode, 0)
            self.assertIn("CLAIM_SOURCE_EXCERPT_NOT_FOUND:0", result.stdout + result.stderr)

    def test_verify_cli_rejects_an_unreadable_primary_tds(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            source = root / "supplier.pdf"
            _write_supplier_pdf(source)
            data = _manifest(source)
            source.write_bytes(b"not a real PDF")
            data["source_files"][1]["sha256"] = hashlib.sha256(source.read_bytes()).hexdigest()
            manifest = root / "sources.yaml"
            manifest.write_text(yaml.safe_dump(data, sort_keys=False), encoding="utf-8")

            result = subprocess.run(
                [
                    str(PYTHON),
                    str(AGENT_ROOT / "scripts" / "verify_tds.py"),
                    "--manifest",
                    str(manifest),
                    "--manifest-only",
                ],
                cwd=REPO_ROOT,
                capture_output=True,
                text=True,
            )

            self.assertNotEqual(result.returncode, 0)
            self.assertIn("PRIMARY_TDS_UNREADABLE:S001", result.stdout + result.stderr)

    def test_verify_cli_rejects_a_mapping_without_both_product_identities(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            source = root / "supplier.pdf"
            _write_supplier_pdf(source)
            data = _manifest(source)
            mapping = root / "product-mapping.txt"
            mapping.write_text("UNRELATED -> OTHER", encoding="utf-8")
            data["source_files"][0]["sha256"] = hashlib.sha256(mapping.read_bytes()).hexdigest()
            manifest = root / "sources.yaml"
            manifest.write_text(yaml.safe_dump(data, sort_keys=False), encoding="utf-8")

            result = subprocess.run(
                [
                    str(PYTHON),
                    str(AGENT_ROOT / "scripts" / "verify_tds.py"),
                    "--manifest",
                    str(manifest),
                    "--manifest-only",
                ],
                cwd=REPO_ROOT,
                capture_output=True,
                text=True,
            )

            self.assertNotEqual(result.returncode, 0)
            self.assertIn("PRODUCT_MAPPING_PAIR_MISMATCH:S000", result.stdout + result.stderr)

    def test_verify_cli_pins_the_template_hash_and_embedded_logo(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            source = root / "supplier.pdf"
            _write_supplier_pdf(source)
            data = _manifest(source)
            data["template"]["sha256"] = "0" * 64
            data["template"]["embedded_logo_source_id"] = "S001"
            manifest = root / "sources.yaml"
            manifest.write_text(yaml.safe_dump(data, sort_keys=False), encoding="utf-8")

            result = subprocess.run(
                [
                    str(PYTHON),
                    str(AGENT_ROOT / "scripts" / "verify_tds.py"),
                    "--manifest",
                    str(manifest),
                    "--manifest-only",
                ],
                cwd=REPO_ROOT,
                capture_output=True,
                text=True,
            )

            combined = result.stdout + result.stderr
            self.assertNotEqual(result.returncode, 0)
            self.assertIn("TEMPLATE_HASH_MISMATCH", combined)
            self.assertIn("TEMPLATE_LOGO_SOURCE_INVALID", combined)

    def test_build_cli_rejects_a_template_other_than_the_pinned_template(self) -> None:
        with _official_tds_directory() as temp_dir:
            root = Path(temp_dir)
            source = root / "supplier.pdf"
            _write_supplier_pdf(source)
            data = _manifest(source)
            manifest = root / "sources.yaml"
            manifest.write_text(yaml.safe_dump(data, sort_keys=False), encoding="utf-8")
            alternate = root / "alternate.docx"
            document = Document(AGENT_ROOT / "templates" / "tiovar-tds-template.docx")
            document.core_properties.title = "Unpinned alternate"
            document.save(alternate)

            result = subprocess.run(
                [
                    str(PYTHON),
                    str(AGENT_ROOT / "scripts" / "build_tds.py"),
                    "--manifest",
                    str(manifest),
                    "--template",
                    str(alternate),
                    "--output-dir",
                    str(root),
                    "--docx-only",
                ],
                cwd=REPO_ROOT,
                capture_output=True,
                text=True,
            )

            self.assertNotEqual(result.returncode, 0)
            self.assertIn("TEMPLATE_ARGUMENT_MISMATCH", result.stdout + result.stderr)

    @unittest.skipUnless(sys.platform == "win32", "Full artifact generation is Windows-specific")
    def test_full_verify_requires_hashes_visual_review_and_manifest_content(self) -> None:
        with _official_tds_directory() as temp_dir:
            root = Path(temp_dir)
            source = root / "supplier.pdf"
            _write_supplier_pdf(source)
            manifest = root / "sources.yaml"
            manifest.write_text(yaml.safe_dump(_manifest(source), sort_keys=False), encoding="utf-8")
            build = subprocess.run(
                [
                    str(PYTHON),
                    str(AGENT_ROOT / "scripts" / "build_tds.py"),
                    "--manifest",
                    str(manifest),
                    "--template",
                    str(AGENT_ROOT / "templates" / "tiovar-tds-template.docx"),
                    "--output-dir",
                    str(root),
                ],
                cwd=REPO_ROOT,
                capture_output=True,
                text=True,
            )
            self.assertEqual(build.returncode, 0, build.stdout + build.stderr)
            data = yaml.safe_load(manifest.read_text(encoding="utf-8"))
            data["artifacts"]["docx_sha256"] = None
            data["artifacts"]["pdf_sha256"] = None
            data["content"]["technical_data"][0]["value"] = "999"
            manifest.write_text(yaml.safe_dump(data, sort_keys=False), encoding="utf-8")

            result = subprocess.run(
                [
                    str(PYTHON),
                    str(AGENT_ROOT / "scripts" / "verify_tds.py"),
                    "--manifest",
                    str(manifest),
                    "--docx",
                    str(root / "tiovar-tp-x100-en.docx"),
                    "--pdf",
                    str(root / "tiovar-tp-x100-en.pdf"),
                ],
                cwd=REPO_ROOT,
                capture_output=True,
                text=True,
            )

            combined = result.stdout + result.stderr
            self.assertNotEqual(result.returncode, 0)
            self.assertIn("ARTIFACT_HASH_REQUIRED:DOCX", combined)
            self.assertIn("ARTIFACT_HASH_REQUIRED:PDF", combined)
            self.assertIn("VISUAL_REVIEW_NOT_PASSED", combined)
            self.assertIn("DOCX_CONTENT_MISSING:technical_data:999", combined)
            self.assertIn("PDF_CONTENT_MISSING:technical_data:999", combined)

    def test_build_cli_rejects_public_output_and_user_approved_status(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            source = root / "supplier.pdf"
            _write_supplier_pdf(source)
            data = _manifest(source)
            data["product"]["status"] = "user-approved"
            manifest = root / "sources.yaml"
            manifest.write_text(yaml.safe_dump(data, sort_keys=False), encoding="utf-8")

            result = subprocess.run(
                [
                    str(PYTHON),
                    str(AGENT_ROOT / "scripts" / "build_tds.py"),
                    "--manifest",
                    str(manifest),
                    "--template",
                    str(AGENT_ROOT / "templates" / "tiovar-tds-template.docx"),
                    "--output-dir",
                    str(root / "public" / "documents" / "tds"),
                    "--docx-only",
                ],
                cwd=REPO_ROOT,
                capture_output=True,
                text=True,
            )

            self.assertNotEqual(result.returncode, 0)
            combined = result.stdout + result.stderr
            self.assertIn("BUILD_STATUS_NOT_DRAFT", combined)
            self.assertIn("OUTPUT_OUTSIDE_MANIFEST_DIRECTORY", combined)

    def test_build_cli_rejects_artifact_path_traversal(self) -> None:
        with _official_tds_directory() as temp_dir:
            root = Path(temp_dir)
            source = root / "supplier.pdf"
            _write_supplier_pdf(source)
            data = _manifest(source)
            escaped = root / "escaped.docx"
            data["artifacts"]["docx"] = str(escaped)
            manifest = root / "sources.yaml"
            manifest.write_text(yaml.safe_dump(data, sort_keys=False), encoding="utf-8")

            result = subprocess.run(
                [
                    str(PYTHON),
                    str(AGENT_ROOT / "scripts" / "build_tds.py"),
                    "--manifest",
                    str(manifest),
                    "--template",
                    str(AGENT_ROOT / "templates" / "tiovar-tds-template.docx"),
                    "--output-dir",
                    str(root),
                    "--docx-only",
                ],
                cwd=REPO_ROOT,
                capture_output=True,
                text=True,
            )

            self.assertNotEqual(result.returncode, 0)
            self.assertIn("ARTIFACT_NAME_INVALID:DOCX", result.stdout + result.stderr)
            self.assertFalse(escaped.exists())

    def test_build_cli_refuses_to_overwrite_an_existing_draft_by_default(self) -> None:
        with _official_tds_directory() as temp_dir:
            root = Path(temp_dir)
            source = root / "supplier.pdf"
            _write_supplier_pdf(source)
            data = _manifest(source)
            manifest = root / "sources.yaml"
            manifest.write_text(yaml.safe_dump(data, sort_keys=False), encoding="utf-8")
            existing = root / "tiovar-tp-x100-en.docx"
            existing.write_bytes(b"user-edited draft")

            result = subprocess.run(
                [
                    str(PYTHON),
                    str(AGENT_ROOT / "scripts" / "build_tds.py"),
                    "--manifest",
                    str(manifest),
                    "--template",
                    str(AGENT_ROOT / "templates" / "tiovar-tds-template.docx"),
                    "--output-dir",
                    str(root),
                    "--docx-only",
                ],
                cwd=REPO_ROOT,
                capture_output=True,
                text=True,
            )

            self.assertNotEqual(result.returncode, 0)
            self.assertIn("OUTPUT_EXISTS:DOCX", result.stdout + result.stderr)
            self.assertEqual(existing.read_bytes(), b"user-edited draft")

    def test_replace_existing_draft_requires_matching_recorded_hash(self) -> None:
        with _official_tds_directory() as temp_dir:
            root = Path(temp_dir)
            source = root / "supplier.pdf"
            _write_supplier_pdf(source)
            data = _manifest(source)
            manifest = root / "sources.yaml"
            manifest.write_text(yaml.safe_dump(data, sort_keys=False), encoding="utf-8")
            existing = root / "tiovar-tp-x100-en.docx"
            existing.write_bytes(b"user-edited draft")

            result = subprocess.run(
                [
                    str(PYTHON),
                    str(AGENT_ROOT / "scripts" / "build_tds.py"),
                    "--manifest",
                    str(manifest),
                    "--template",
                    str(AGENT_ROOT / "templates" / "tiovar-tds-template.docx"),
                    "--output-dir",
                    str(root),
                    "--docx-only",
                    "--replace-existing-draft",
                ],
                cwd=REPO_ROOT,
                capture_output=True,
                text=True,
            )

            self.assertNotEqual(result.returncode, 0)
            self.assertIn("EXISTING_ARTIFACT_HASH_REQUIRED:DOCX", result.stdout + result.stderr)
            self.assertEqual(existing.read_bytes(), b"user-edited draft")

    @unittest.skipUnless(sys.platform == "win32", "Renderer failure path is Windows-specific")
    def test_render_failure_does_not_partially_overwrite_existing_draft(self) -> None:
        with _official_tds_directory() as temp_dir:
            root = Path(temp_dir)
            source = root / "supplier.pdf"
            _write_supplier_pdf(source)
            data = _manifest(source)
            existing = root / "tiovar-tp-x100-en.docx"
            existing.write_bytes((AGENT_ROOT / "templates" / "tiovar-tds-template.docx").read_bytes())
            original_hash = hashlib.sha256(existing.read_bytes()).hexdigest()
            data["artifacts"]["docx_sha256"] = original_hash
            manifest = root / "sources.yaml"
            manifest.write_text(yaml.safe_dump(data, sort_keys=False), encoding="utf-8")
            env = dict(os.environ)
            env["PATH"] = ""

            result = subprocess.run(
                [
                    str(PYTHON),
                    str(AGENT_ROOT / "scripts" / "build_tds.py"),
                    "--manifest",
                    str(manifest),
                    "--template",
                    str(AGENT_ROOT / "templates" / "tiovar-tds-template.docx"),
                    "--output-dir",
                    str(root),
                    "--replace-existing-draft",
                ],
                cwd=REPO_ROOT,
                env=env,
                capture_output=True,
                text=True,
            )

            self.assertNotEqual(result.returncode, 0)
            self.assertIn("POWERSHELL_NOT_FOUND", result.stdout + result.stderr)
            self.assertEqual(hashlib.sha256(existing.read_bytes()).hexdigest(), original_hash)

    def test_build_cli_rejects_manifest_outside_official_tds_root(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            source = root / "supplier.pdf"
            _write_supplier_pdf(source)
            manifest = root / "sources.yaml"
            manifest.write_text(yaml.safe_dump(_manifest(source), sort_keys=False), encoding="utf-8")

            result = subprocess.run(
                [
                    str(PYTHON),
                    str(AGENT_ROOT / "scripts" / "build_tds.py"),
                    "--manifest",
                    str(manifest),
                    "--template",
                    str(AGENT_ROOT / "templates" / "tiovar-tds-template.docx"),
                    "--output-dir",
                    str(root),
                    "--docx-only",
                ],
                cwd=REPO_ROOT,
                capture_output=True,
                text=True,
            )

            self.assertNotEqual(result.returncode, 0)
            self.assertIn("OUTPUT_OUTSIDE_TDS_ROOT", result.stdout + result.stderr)
            self.assertFalse((root / "tiovar-tp-x100-en.docx").exists())

    def test_mapping_requires_a_single_paired_identity_record(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            source = root / "supplier.pdf"
            _write_supplier_pdf(source)
            data = _manifest(source)
            mapping = root / "product-mapping.txt"
            mapping.write_text("TP-X100 -> OTHER\nTP-Y200 -> SUP-100", encoding="utf-8")
            data["source_files"][0]["sha256"] = hashlib.sha256(mapping.read_bytes()).hexdigest()
            manifest = root / "sources.yaml"
            manifest.write_text(yaml.safe_dump(data, sort_keys=False), encoding="utf-8")

            result = subprocess.run(
                [str(PYTHON), str(AGENT_ROOT / "scripts" / "verify_tds.py"), "--manifest", str(manifest), "--manifest-only"],
                cwd=REPO_ROOT,
                capture_output=True,
                text=True,
            )

            self.assertNotEqual(result.returncode, 0)
            self.assertIn("PRODUCT_MAPPING_PAIR_MISMATCH:S000", result.stdout + result.stderr)

    def test_technical_source_requires_property_and_exact_value_pair(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            source = root / "supplier.pdf"
            _write_supplier_pdf(source)
            data = _manifest(source)
            data["content"]["technical_data"][0]["property"] = "Invented opacity"
            data["content"]["technical_data"][0]["source_locator"] = "page 1 / specifications / Invented opacity"
            manifest = root / "sources.yaml"
            manifest.write_text(yaml.safe_dump(data, sort_keys=False), encoding="utf-8")

            result = subprocess.run(
                [str(PYTHON), str(AGENT_ROOT / "scripts" / "verify_tds.py"), "--manifest", str(manifest), "--manifest-only"],
                cwd=REPO_ROOT,
                capture_output=True,
                text=True,
            )

            self.assertNotEqual(result.returncode, 0)
            self.assertIn("TECHNICAL_PROPERTY_VALUE_PAIR_NOT_FOUND:0", result.stdout + result.stderr)

    def test_all_customer_visible_fields_require_claim_mapping(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            source = root / "supplier.pdf"
            _write_supplier_pdf(source)
            data = _manifest(source)
            data["claims"] = [claim for claim in data["claims"] if claim["field"] != "packaging_note"]
            manifest = root / "sources.yaml"
            manifest.write_text(yaml.safe_dump(data, sort_keys=False), encoding="utf-8")

            result = subprocess.run(
                [str(PYTHON), str(AGENT_ROOT / "scripts" / "verify_tds.py"), "--manifest", str(manifest), "--manifest-only"],
                cwd=REPO_ROOT,
                capture_output=True,
                text=True,
            )

            self.assertNotEqual(result.returncode, 0)
            self.assertIn("CLAIM_SOURCE_MISSING:packaging_note:Packaging is not confirmed.", result.stdout + result.stderr)

    @unittest.skipUnless(sys.platform == "win32", "Full artifact generation is Windows-specific")
    def test_full_verify_rejects_unregistered_artifact_text_and_wrong_status_label(self) -> None:
        with _official_tds_directory() as temp_dir:
            root = Path(temp_dir)
            source = root / "supplier.pdf"
            _write_supplier_pdf(source)
            manifest = root / "sources.yaml"
            manifest.write_text(yaml.safe_dump(_manifest(source), sort_keys=False), encoding="utf-8")
            build = subprocess.run(
                [
                    str(PYTHON), str(AGENT_ROOT / "scripts" / "build_tds.py"),
                    "--manifest", str(manifest),
                    "--template", str(AGENT_ROOT / "templates" / "tiovar-tds-template.docx"),
                    "--output-dir", str(root),
                ],
                cwd=REPO_ROOT,
                capture_output=True,
                text=True,
            )
            self.assertEqual(build.returncode, 0, build.stdout + build.stderr)
            docx = root / "tiovar-tp-x100-en.docx"
            pdf = root / "tiovar-tp-x100-en.pdf"
            document = Document(docx)
            document.add_paragraph("USER-APPROVED UNREGISTERED PERFORMANCE CLAIM")
            document.save(docx)
            data = yaml.safe_load(manifest.read_text(encoding="utf-8"))
            data["artifacts"]["docx_sha256"] = hashlib.sha256(docx.read_bytes()).hexdigest()
            data["qa"]["visual_review"] = {
                "status": "passed",
                "page_count": 1,
                "docx_sha256": data["artifacts"]["docx_sha256"],
                "pdf_sha256": data["artifacts"]["pdf_sha256"],
                "reviewed_at": "2026-08-26",
            }
            manifest.write_text(yaml.safe_dump(data, sort_keys=False), encoding="utf-8")

            result = subprocess.run(
                [
                    str(PYTHON), str(AGENT_ROOT / "scripts" / "verify_tds.py"),
                    "--manifest", str(manifest), "--docx", str(docx), "--pdf", str(pdf),
                ],
                cwd=REPO_ROOT,
                capture_output=True,
                text=True,
            )

            self.assertNotEqual(result.returncode, 0)
            self.assertIn("DOCX_UNEXPECTED_CONTENT", result.stdout + result.stderr)

    def test_project_instructions_define_the_tds_agent_exception(self) -> None:
        instructions = (REPO_ROOT / "AGENTS.md").read_text(encoding="utf-8")
        self.assertIn("tiovar-tds-agent", instructions)
        self.assertIn("must not create worktrees", instructions)
        self.assertIn("write to `public/`", instructions)


if __name__ == "__main__":
    unittest.main()
